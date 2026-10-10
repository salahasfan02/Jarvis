import asyncio
import zipfile
from pathlib import Path

import docx

from app import channels
from app.tools.builtin import documents


def test_markdown_becomes_a_real_word_document(tmp_path):
    md = """# E-bike research

Intro with **bold** and a [source](https://example.com/a).

## Top picks
- Lectric XP — **$999**
  - folds flat
1. Buy from the maker
2. Check the warranty

| Model | Price |
|-------|-------|
| Lectric | $999 |
| Aventon | $1,399 |

```
range = battery_wh / wh_per_mile
```
> Quote this.
"""
    path = tmp_path / "out.docx"
    words = documents.markdown_to_docx("E-bike research", md, path)
    assert words > 10

    d = docx.Document(str(path))
    paras = [(p.style.name, p.text) for p in d.paragraphs]
    texts = [t for _, t in paras]
    assert texts.count("E-bike research") == 1          # title not repeated
    assert ("Heading 2", "Top picks") in paras
    assert ("List Bullet", "Lectric XP — $999") in paras
    assert ("List Bullet 2", "folds flat") in paras
    assert ("List Number", "Buy from the maker") in paras
    assert any("range = battery_wh" in t for t in texts)
    assert d.tables[0].cell(2, 1).text == "$1,399"
    assert d.tables[0].cell(0, 0).paragraphs[0].runs[0].bold
    rels = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    assert rels == ["https://example.com/a"]
    assert any(r.bold and r.text == "bold" for p in d.paragraphs for r in p.runs)


def test_create_word_document_never_overwrites():
    first = documents.create_word_document("Weekly notes", "hello")
    second = documents.create_word_document("Weekly notes", "hello again")
    p1 = Path(first.split(" at ")[1].split(". ")[0])
    p2 = Path(second.split(" at ")[1].split(". ")[0])
    assert p1 != p2 and p1.exists() and p2.exists()
    assert p1.parent == Path.home() / "Documents" / "Jarvis"


class FakeChannel(channels.ChannelSession):
    name = "Fake"

    def __init__(self):
        self.sent = []

    async def send_file(self, path, caption=""):
        self.sent.append((path.name, path.read_bytes(), caption))
        return f"sent {path.name}"


def _with_channel(channel, coro):
    async def run():
        token = channels.current.set(channel)
        try:
            return await coro
        finally:
            channels.current.reset(token)
    return asyncio.run(run())


def test_send_file_delivers_through_the_active_channel(tmp_path):
    f = tmp_path / "lease.pdf"
    f.write_bytes(b"%PDF-1.4 lease")
    ch = FakeChannel()
    assert _with_channel(ch, documents.send_file(str(f), "your lease")) == "sent lease.pdf"
    assert ch.sent == [("lease.pdf", b"%PDF-1.4 lease", "your lease")]


def test_send_file_zips_folders(tmp_path):
    folder = tmp_path / "project"
    (folder / "src").mkdir(parents=True)
    (folder / "src" / "main.py").write_text("print('hi')")
    ch = FakeChannel()
    _with_channel(ch, documents.send_file(str(folder)))
    name, data, _ = ch.sent[0]
    assert name == "project.zip"
    archive = tmp_path / "check.zip"
    archive.write_bytes(data)
    assert zipfile.ZipFile(archive).namelist() == ["src/", "src/main.py"]


def test_send_file_refuses_oversized_and_missing_files(tmp_path):
    big = tmp_path / "big.bin"
    big.write_bytes(b"x" * 2048)
    ch = FakeChannel()
    ch.max_file_bytes = 1024
    assert "over the" in _with_channel(ch, documents.send_file(str(big)))
    assert "No file or folder" in _with_channel(ch, documents.send_file(str(tmp_path / "nope")))
    assert ch.sent == []


def test_send_file_at_the_mac_does_not_send(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("x")
    assert "nothing was sent" in asyncio.run(documents.send_file(str(f)))
