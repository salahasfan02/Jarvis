"""Spreadsheets, slide decks, reading/comparing documents, PDF forms."""
import asyncio
from pathlib import Path

from app.tools.builtin import office


def test_spreadsheet_keeps_numbers_and_header(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "DOCS_DIR", tmp_path)
    out = office.create_spreadsheet("Budget", rows=[["Item", "Cost"], ["Rent", "1,200"],
                                                     ["Food", "250.5"], ["Total", "=SUM(B2:B3)"]])
    path = Path(out.split(" at ")[1].split(". To")[0])
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    assert ws["A1"].value == "Item" and ws["B2"].value == 1200 and ws["B3"].value == 250.5
    assert ws["B4"].value == "=SUM(B2:B3)"


def test_spreadsheet_from_csv(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "DOCS_DIR", tmp_path)
    assert "2 rows" in office.create_spreadsheet("People", csv="name,age\nSam,30\nAli,25")


def test_presentation_round_trip(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "DOCS_DIR", tmp_path)
    out = office.create_presentation("Q3 Review", [
        {"title": "Wins", "bullets": ["Revenue up 12%", "Two new clients"], "notes": "smile"},
        {"title": "Next", "bullets": "- Hire\n- Ship v2"}])
    path = Path(out.split(" at ")[1].rstrip("."))
    text = office.extract(path)
    assert "Revenue up 12%" in text and "Ship v2" in text and "[slide 3]" in text


def test_read_and_compare_documents(tmp_path, monkeypatch):
    a, b = tmp_path / "v1.txt", tmp_path / "v2.txt"
    a.write_text("Rent is 1000 pounds. The term is one year. Pets are allowed.")
    b.write_text("Rent is 1200 pounds. The term is one year. Pets are not allowed.")
    monkeypatch.setattr(office, "find_file", lambda q: (Path(q), []))
    read = asyncio.run(office.read_document(str(a)))
    assert "Rent is 1000" in read
    diff = asyncio.run(office.compare_documents(str(a), str(b)))
    assert "- Rent is 1000 pounds." in diff and "+ Rent is 1200 pounds." in diff


def test_pdf_form_lists_then_fills(tmp_path, monkeypatch):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (ArrayObject, DictionaryObject, FloatObject, NameObject,
                               NumberObject, TextStringObject)
    w = PdfWriter()
    page = w.add_blank_page(300, 300)
    field = DictionaryObject({
        NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("full_name"),
        NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
        NameObject("/Rect"): ArrayObject([FloatObject(10), FloatObject(10), FloatObject(200),
                                          FloatObject(40)]),
        NameObject("/F"): NumberObject(4)})
    ref = w._add_object(field)
    page[NameObject("/Annots")] = ArrayObject([ref])
    w._root_object[NameObject("/AcroForm")] = DictionaryObject(
        {NameObject("/Fields"): ArrayObject([ref])})
    form = tmp_path / "form.pdf"
    with open(form, "wb") as fh:
        w.write(fh)
    monkeypatch.setattr(office, "find_file", lambda q: (Path(q), []))
    monkeypatch.setattr(office, "DOCS_DIR", tmp_path)
    assert "full_name" in asyncio.run(office.fill_pdf_form(str(form)))
    out = asyncio.run(office.fill_pdf_form(str(form), {"full_name": "Salah Asfan"}))
    filled = Path(out.split("saved ")[1].split(".pdf")[0] + ".pdf")
    assert PdfReader(str(filled)).get_fields()["full_name"]["/V"] == "Salah Asfan"
