"""Document work: read and compare PDFs / Word / Excel / PowerPoint, create
spreadsheets and slide decks, fill in PDF forms, and index a folder so the
user's documents can be searched by meaning (search_documents)."""
from __future__ import annotations

import asyncio
import csv
import difflib
import io
import re
import time
from pathlib import Path

from ..base import tool
from ... import activity
from .documents import DOCS_DIR, find_file

READ_LIMIT = 12000
INDEX_EXT = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".pptx", ".rtf", ".csv"}


def _resolve(name: str) -> tuple[Path | None, str]:
    p, candidates = find_file(name)
    if p is None:
        if candidates:
            return None, ("Several files could match — ask which, then pass the full path:\n"
                          + "\n".join(candidates))
        return None, f"No file called '{name}' was found."
    if p.is_dir():
        return None, f"{p} is a folder, not a document."
    return p, ""


def extract(path: Path, pages: str = "") -> str:
    """Plain text of a document. `pages` like '1-3,7' limits a PDF."""
    suffix = path.suffix.lower()
    data = path.read_bytes()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        wanted = _page_set(pages, len(reader.pages))
        return "\n\n".join(f"[page {i + 1}]\n{(reader.pages[i].extract_text() or '').strip()}"
                           for i in wanted)
    if suffix == ".pptx":
        from pptx import Presentation
        deck = Presentation(io.BytesIO(data))
        out = []
        for n, slide in enumerate(deck.slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            out.append(f"[slide {n}]\n" + "\n".join(texts))
        return "\n\n".join(out)
    from ...knowledge import attachments
    if suffix == ".docx":
        return attachments._extract_docx(data)
    if suffix in (".xlsx", ".xlsm"):
        return attachments._extract_xlsx(data)
    return data.decode("utf-8", errors="replace")


def _page_set(spec: str, total: int) -> list[int]:
    if not spec.strip():
        return list(range(total))
    pages: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            pages += range(int(a) - 1, min(int(b), total))
        elif part.isdigit():
            pages.append(int(part) - 1)
    return [p for p in dict.fromkeys(pages) if 0 <= p < total]


@tool(
    name="read_document",
    description="Read the text of a PDF, Word, Excel, PowerPoint or text file for YOURSELF, "
                "to summarize it, answer questions about it or quote it. Give a path or "
                "just the name (\"the lease\", \"my CV\"). For long PDFs pass pages, e.g. "
                "'1-5'. To show a file to the user (\"open the lease\"), use open_file instead.",
    parameters={"type": "object", "properties": {
        "path": {"type": "string", "description": "file path or name"},
        "pages": {"type": "string", "description": "PDF pages, e.g. '1-3,8' (optional)"},
    }, "required": ["path"]},
    agent_tags=["files", "research", "messaging"],
)
async def read_document(path: str, pages: str = "") -> str:
    p, err = await asyncio.to_thread(_resolve, path)
    if err:
        return err
    try:
        text = await asyncio.to_thread(extract, p, pages)
    except Exception as e:
        return f"Couldn't read {p.name}: {e}"
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return f"{p.name} has no readable text (it may be a scanned image — try screen/OCR)."
    more = (f"\n\n… ({len(text) - READ_LIMIT} more characters — pass pages to read further)"
            if len(text) > READ_LIMIT else "")
    return f"{p} ({len(text)} characters):\n\n{text[:READ_LIMIT]}{more}"


@tool(
    name="compare_documents",
    description="Compare two documents (PDF/Word/text): returns what was added, removed or "
                "changed between them, e.g. two versions of a contract or CV.",
    parameters={"type": "object", "properties": {
        "first": {"type": "string", "description": "path or name of the older/first file"},
        "second": {"type": "string", "description": "path or name of the newer/second file"},
    }, "required": ["first", "second"]},
    agent_tags=["files", "research"],
)
async def compare_documents(first: str, second: str) -> str:
    texts, paths = [], []
    for name in (first, second):
        p, err = await asyncio.to_thread(_resolve, name)
        if err:
            return err
        paths.append(p)
        texts.append(await asyncio.to_thread(extract, p))

    def lines(t: str) -> list[str]:
        return [ln.strip() for ln in re.split(r"(?<=[.!?])\s+|\n", t) if ln.strip()]
    a, b = lines(texts[0]), lines(texts[1])
    ratio = difflib.SequenceMatcher(None, a, b).ratio()
    diff = [f"{d[0]} {d[1:]}" for d in difflib.unified_diff(a, b, lineterm="", n=0)
            if d.startswith(("+", "-")) and not d.startswith(("+++", "---"))]
    shown = "\n".join(diff[:150])
    more = f"\n… {len(diff) - 150} more changed lines" if len(diff) > 150 else ""
    return (f"{paths[0].name} vs {paths[1].name}: {round(ratio * 100)}% the same, "
            f"{sum(d.startswith('-') for d in diff)} lines removed, "
            f"{sum(d.startswith('+') for d in diff)} added.\n"
            f"(- only in the first, + only in the second)\n{shown}{more}")


def _out_path(filename: str, title: str, ext: str) -> Path:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    stem = Path(filename).stem if filename else re.sub(r"[^\w\- ]+", "", title).strip() or "document"
    stem = re.sub(r"\s+", "-", stem.lower())[:80]
    path, n = DOCS_DIR / f"{stem}{ext}", 2
    while path.exists():
        path, n = DOCS_DIR / f"{stem}-{n}{ext}", n + 1
    return path


@tool(
    name="create_spreadsheet",
    description="Create an Excel spreadsheet (.xlsx) from rows of data — a budget, a list, "
                "a table of results. First row = column headers. Numbers stay numbers; "
                "formulas like '=SUM(B2:B9)' work. Saved in ~/Documents/Jarvis; returns "
                "the path (attach it to an email or send_file it).",
    parameters={"type": "object", "properties": {
        "title": {"type": "string"},
        "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}},
                 "description": "list of rows, each a list of cells; first row is the header"},
        "csv": {"type": "string", "description": "alternatively, the data as CSV text"},
        "filename": {"type": "string", "description": "optional file name"},
    }, "required": ["title"]},
)
def create_spreadsheet(title: str, rows: list | None = None, csv_text: str = "",
                       filename: str = "", **kw) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    csv_text = csv_text or kw.get("csv", "")
    if not rows and csv_text:
        rows = list(csv.reader(io.StringIO(csv_text.strip())))
    if not rows:
        return "No data given — pass rows (first row = headers) or csv."

    def cell(v):
        if isinstance(v, str):
            s = v.strip()
            if re.fullmatch(r"-?[\d,]*\.?\d+", s) and not re.fullmatch(r"0\d+", s):
                try:
                    return float(s.replace(",", "")) if "." in s else int(s.replace(",", ""))
                except ValueError:
                    return s
            return s
        return v
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Sheet1"
    for r in rows:
        ws.append([cell(v) for v in (r if isinstance(r, list) else [r])])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6F43")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    path = _out_path(filename, title, ".xlsx")
    wb.save(path)
    activity.offer_undo(kind="remove_created", path=str(path))
    return (f"Created spreadsheet '{title}' ({len(rows) - 1} rows) at {path}. To email it, "
            "attach this path; to send it to the phone, send_file it.")


@tool(
    name="create_presentation",
    description="Create a PowerPoint slide deck (.pptx): a title slide plus content slides, "
                "each with a title and bullet points (and optional speaker notes). Saved in "
                "~/Documents/Jarvis; returns the path.",
    parameters={"type": "object", "properties": {
        "title": {"type": "string"},
        "subtitle": {"type": "string"},
        "slides": {"type": "array", "items": {"type": "object", "properties": {
            "title": {"type": "string"},
            "bullets": {"type": "array", "items": {"type": "string"}},
            "notes": {"type": "string"}}}},
        "filename": {"type": "string"},
    }, "required": ["title", "slides"]},
)
def create_presentation(title: str, slides: list, subtitle: str = "", filename: str = "") -> str:
    from pptx import Presentation
    from pptx.util import Pt
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = title
    if len(first.placeholders) > 1:
        first.placeholders[1].text = subtitle or time.strftime("%B %Y")
    for s in slides or []:
        s = s if isinstance(s, dict) else {"title": str(s), "bullets": []}
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        body = slide.placeholders[1].text_frame
        bullets = s.get("bullets") or []
        if isinstance(bullets, str):
            bullets = [b.strip("-• ") for b in bullets.splitlines() if b.strip()]
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(20 if len(bullets) <= 5 else 16)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    path = _out_path(filename, title, ".pptx")
    deck.save(path)
    activity.offer_undo(kind="remove_created", path=str(path))
    return f"Created presentation '{title}' ({len(slides or [])} slides + title) at {path}."


@tool(
    name="fill_pdf_form",
    description="Fill in a fillable PDF form. Call with only `path` first to list the form's "
                "fields; then call again with `values` {field name: value}. Saves a filled "
                "copy next to Jarvis's documents (the original is untouched).",
    parameters={"type": "object", "properties": {
        "path": {"type": "string", "description": "the PDF form (path or name)"},
        "values": {"type": "object", "description": "field name -> value"},
    }, "required": ["path"]},
    risk="confirm",
    agent_tags=["files"],
)
async def fill_pdf_form(path: str, values: dict | None = None) -> str:
    from pypdf import PdfReader, PdfWriter
    p, err = await asyncio.to_thread(_resolve, path)
    if err:
        return err
    reader = PdfReader(str(p))
    fields = reader.get_fields() or {}
    if not fields:
        return f"{p.name} has no fillable fields (it may need to be filled by hand or printed)."
    if not values:
        lines = []
        for name, f in fields.items():
            current = f.get("/V")
            opts = f.get("/_States_") or f.get("/Opt")
            lines.append(f"- {name}" + (f" = {current}" if current else "")
                         + (f" (options: {', '.join(map(str, opts))})" if opts else ""))
        return f"{p.name} has {len(fields)} fields:\n" + "\n".join(lines)
    unknown = [k for k in values if k not in fields]
    writer = PdfWriter()
    writer.append(reader)
    for page in writer.pages:
        writer.update_page_form_field_values(page, {k: str(v) for k, v in values.items()
                                                    if k in fields}, auto_regenerate=False)
    writer.set_need_appearances_writer(True)
    out = _out_path("", f"{p.stem} filled", ".pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    activity.offer_undo(kind="remove_created", path=str(out))
    note = f" Unknown fields skipped: {', '.join(unknown)}." if unknown else ""
    return f"Filled {len(values) - len(unknown)} fields and saved {out}.{note}"


_index_state = {"running": False, "done": 0, "total": 0, "folder": ""}


async def _index(folder: Path, files: list[Path]) -> None:
    from ... import db
    from ...knowledge import store
    existing = {d["name"]: d for d in db.list_documents()}
    try:
        for f in files:
            key = str(f)
            old = existing.get(key)
            if old and old["created_at"] >= f.stat().st_mtime:
                _index_state["done"] += 1
                continue
            try:
                text = await asyncio.to_thread(extract, f)
                if old:
                    db.delete_document(old["id"])
                chunks = store.chunk_text(text)[:200]
                if chunks:
                    from ...llm import ollama_client
                    embedded = [(c, await ollama_client.embed(c)) for c in chunks]
                    db.add_document(key, embedded)
            except Exception:
                pass
            _index_state["done"] += 1
    finally:
        _index_state["running"] = False


@tool(
    name="index_documents",
    description="Make a folder's documents (PDF, Word, Excel, PowerPoint, text) searchable "
                "by meaning with search_documents — e.g. index ~/Documents once, then "
                "\"find the contract that mentions a break clause\". Runs in the background; "
                "already-indexed, unchanged files are skipped. Call with no folder to see progress.",
    parameters={"type": "object", "properties": {
        "folder": {"type": "string", "description": "folder to index, e.g. ~/Documents"},
    }},
    agent_tags=["files", "research"],
)
async def index_documents(folder: str = "") -> str:
    if not folder:
        s = _index_state
        if s["running"]:
            return f"Indexing {s['folder']}: {s['done']}/{s['total']} files done."
        return (f"Last indexing of {s['folder'] or '(nothing yet)'} finished: {s['done']} files."
                if s["folder"] else "Nothing has been indexed yet — pass a folder.")
    if _index_state["running"]:
        return f"Already indexing {_index_state['folder']} ({_index_state['done']}/{_index_state['total']})."
    root = Path(folder).expanduser()
    if not root.is_dir():
        return f"Not a folder: {root}"
    skip = {"node_modules", ".git", ".venv", "venv", "Library", "__pycache__"}
    files = [f for f in root.rglob("*")
             if f.is_file() and f.suffix.lower() in INDEX_EXT and f.stat().st_size < 25_000_000
             and not any(part in skip or part.startswith(".") for part in f.relative_to(root).parts)]
    files = sorted(files, key=lambda f: f.stat().st_mtime, reverse=True)[:1500]
    _index_state.update(running=True, done=0, total=len(files), folder=str(root))
    asyncio.get_running_loop().create_task(_index(root, files))
    return (f"Indexing {len(files)} documents in {root} in the background. Searchable with "
            "search_documents as they finish; ask me for progress any time.")
